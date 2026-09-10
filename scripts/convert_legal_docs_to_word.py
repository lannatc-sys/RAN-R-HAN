import os
import re
import sys
import docx

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

FONT_PRIMARY = 'TH Sarabun PSK'
FONT_SECONDARY = 'Calibri'

COLOR_TITLE = RGBColor(0x1B, 0x36, 0x5D)      # Deep Navy
COLOR_H1 = RGBColor(0x1E, 0x3A, 0x8A)         # Royal Navy
COLOR_H2 = RGBColor(0x25, 0x63, 0xEB)         # Bright Blue
COLOR_H3 = RGBColor(0x1F, 0x29, 0x37)         # Charcoal Dark Gray
COLOR_BODY = RGBColor(0x33, 0x41, 0x55)       # Slate Gray
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)      # Muted Gray
COLOR_CODE = RGBColor(0xB4, 0x53, 0x09)       # Amber-700
COLOR_HIGHLIGHT = RGBColor(0xB9, 0x1C, 0x1C)  # Red Alert

def set_run_font(run, font_name=FONT_PRIMARY, size_pt=14, bold=False, italic=False, color=COLOR_BODY):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    
    # Force Asian & Complex Script font for Thai rendering
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)

def set_cell_background(cell, fill_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top='E2E8F0', bottom='E2E8F0', left='none', right='none'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    if top != 'none':
        borders_xml += f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
    else:
        borders_xml += '<w:top w:val="none"/>'
    if bottom != 'none':
        borders_xml += f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
    else:
        borders_xml += '<w:bottom w:val="none"/>'
    if left != 'none':
        borders_xml += f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{left}"/>'
    else:
        borders_xml += '<w:left w:val="none"/>'
    if right != 'none':
        borders_xml += f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{right}"/>'
    else:
        borders_xml += '<w:right w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def add_callout_box(doc, text_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    # Check if there is red or amber alert in text
    full_text = " ".join(text_lines)
    if '🔴' in full_text:
        bg_color = 'FEF2F2'  # Light red
        border_color = 'EF4444' # Red
    elif '🟠' in full_text or 'สถานะ: ฉบับร่าง' in full_text:
        bg_color = 'FFFBEB'  # Light amber
        border_color = 'F59E0B' # Amber
    else:
        bg_color = 'F8FAFC'  # Light slate
        border_color = '3B82F6' # Blue
        
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)
    
    # Custom left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    for idx, line in enumerate(text_lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, line, base_size=13, default_color=COLOR_BODY)

def add_formatted_text(paragraph, text, base_size=14, default_color=COLOR_BODY):
    # Regex to extract bold (**...**), italic (*...*), code (`...`), link ([...](...))
    pattern = re.compile(r'(\*\*(.*?)\*\*|\*(.*?)\*|`(.*?)`|\[(.*?)\]\((.*?)\))')
    
    last_idx = 0
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            plain_part = text[last_idx:start]
            r = paragraph.add_run(plain_part)
            set_run_font(r, size_pt=base_size, color=default_color)
            
        full = match.group(0)
        if full.startswith('**') and full.endswith('**'):
            content = match.group(2)
            r = paragraph.add_run(content)
            # If contains alert icon or [ต้องกรอก], highlight color
            c = COLOR_HIGHLIGHT if ('[ต้องกรอก]' in content or '🔴' in content) else COLOR_TITLE
            set_run_font(r, size_pt=base_size, bold=True, color=c)
        elif full.startswith('*') and full.endswith('*'):
            content = match.group(3)
            r = paragraph.add_run(content)
            set_run_font(r, size_pt=base_size, italic=True, color=default_color)
        elif full.startswith('`') and full.endswith('`'):
            content = match.group(4)
            r = paragraph.add_run(content)
            set_run_font(r, font_name=FONT_SECONDARY, size_pt=base_size-1, bold=True, color=COLOR_CODE)
        elif full.startswith('['):
            content = match.group(5)
            r = paragraph.add_run(content)
            set_run_font(r, size_pt=base_size, color=COLOR_H2)
            r.underline = True
            
        last_idx = end
        
    if last_idx < len(text):
        remaining = text[last_idx:]
        r = paragraph.add_run(remaining)
        set_run_font(r, size_pt=base_size, color=default_color)

def parse_markdown_to_docx(md_content, doc, is_first_doc=True):
    lines = md_content.split('\n')
    i = 0
    in_blockquote = False
    blockquote_lines = []
    
    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()
        
        # Check if currently accumulating blockquote
        if trimmed.startswith('>'):
            content = trimmed.lstrip('>').strip()
            blockquote_lines.append(content)
            i += 1
            continue
        elif blockquote_lines:
            add_callout_box(doc, blockquote_lines)
            blockquote_lines = []
            
        # Empty line
        if not trimmed:
            i += 1
            continue
            
        # Horizontal Rule
        if trimmed in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue
            
        # Headings
        if trimmed.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, trimmed[2:], base_size=20, default_color=COLOR_TITLE)
            for r in p.runs:
                r.bold = True
            i += 1
            continue
        elif trimmed.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, trimmed[3:], base_size=16, default_color=COLOR_H1)
            for r in p.runs:
                r.bold = True
            i += 1
            continue
        elif trimmed.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, trimmed[4:], base_size=14.5, default_color=COLOR_H2)
            for r in p.runs:
                r.bold = True
            i += 1
            continue
        elif trimmed.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, trimmed[5:], base_size=13.5, default_color=COLOR_H3)
            for r in p.runs:
                r.bold = True
            i += 1
            continue
            
        # Table
        if trimmed.startswith('|') and '|' in trimmed[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and '|' in lines[i].strip()[1:]:
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                # Parse headers
                header_raw = table_lines[0]
                headers = [c.strip() for c in header_raw.strip('|').split('|')]
                
                # Check if second line is separator
                sep_raw = table_lines[1]
                data_start = 2 if re.match(r'^[\|\s\-\:]+$', sep_raw) else 1
                
                rows_data = []
                for row_line in table_lines[data_start:]:
                    cells = [c.strip() for c in row_line.strip('|').split('|')]
                    # Match column length
                    if len(cells) < len(headers):
                        cells.extend([''] * (len(headers) - len(cells)))
                    elif len(cells) > len(headers):
                        cells = cells[:len(headers)]
                    rows_data.append(cells)
                    
                table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                # Style Header Row
                hdr_cells = table.rows[0].cells
                for col_idx, h_text in enumerate(headers):
                    hdr_cells[col_idx].text = ""
                    set_cell_background(hdr_cells[col_idx], '1E3A8A') # Dark Navy
                    set_cell_margins(hdr_cells[col_idx], top=140, bottom=140, left=160, right=160)
                    set_cell_borders(hdr_cells[col_idx], top='1E3A8A', bottom='0F172A', left='2563EB', right='2563EB')
                    p = hdr_cells[col_idx].paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    add_formatted_text(p, h_text, base_size=13, default_color=RGBColor(0xFF, 0xFF, 0xFF))
                    for r in p.runs:
                        r.bold = True
                        
                # Style Data Rows
                for row_idx, r_data in enumerate(rows_data):
                    row_cells = table.rows[row_idx + 1].cells
                    row_bg = 'FFFFFF' if (row_idx % 2 == 0) else 'F8FAFC' # Alternating rows
                    for col_idx, cell_text in enumerate(r_data):
                        row_cells[col_idx].text = ""
                        set_cell_background(row_cells[col_idx], row_bg)
                        set_cell_margins(row_cells[col_idx], top=120, bottom=120, left=150, right=150)
                        set_cell_borders(row_cells[col_idx], top='E2E8F0', bottom='CBD5E1', left='F1F5F9', right='F1F5F9')
                        p = row_cells[col_idx].paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15
                        add_formatted_text(p, cell_text, base_size=12.5, default_color=COLOR_BODY)
                        
                # Space after table
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_before = Pt(2)
                spacer.paragraph_format.space_after = Pt(4)
            continue
            
        # Bullet List (- or *)
        if re.match(r'^[\-\*]\s+', trimmed):
            bullet_text = re.sub(r'^[\-\*]\s+', '', trimmed)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.3)
            add_formatted_text(p, bullet_text, base_size=13.5, default_color=COLOR_BODY)
            i += 1
            continue
            
        # Numbered List (1. 2. etc)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', trimmed)
        if num_match:
            item_num = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.3)
            add_formatted_text(p, item_text, base_size=13.5, default_color=COLOR_BODY)
            i += 1
            continue
            
        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, trimmed, base_size=13.5, default_color=COLOR_BODY)
        i += 1
        
    # Flush trailing blockquote if any
    if blockquote_lines:
        add_callout_box(doc, blockquote_lines)

def setup_document():
    doc = Document()
    
    # Page Setup: A4, 1 inch margins
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)   # A4 Width
        section.page_height = Inches(11.69) # A4 Height
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("RAN-R-HAN | เอกสารทางกฎหมายและนโยบายความเป็นส่วนตัว (DRAFT)")
        set_run_font(hrun, font_name=FONT_PRIMARY, size_pt=10, italic=True, color=COLOR_MUTED)
        
        # Footer setup
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("เอกสารฉบับร่าง — ต้องผ่านการตรวจสอบโดยที่ปรึกษากฎหมายก่อนนำไปใช้งานจริง")
        set_run_font(frun, font_name=FONT_PRIMARY, size_pt=9.5, italic=True, color=COLOR_MUTED)
        
    return doc

def main():
    base_dir = r"d:\system make\Ran-R-HAN\docs\Privacy Policy & Terms of Service"
    
    files_to_convert = [
        {
            "md_file": "RANRHANข้อกำหนดการใช้งานDRAFT.md",
            "out_docx": "RANRHAN_ข้อกำหนดการใช้งาน_ToS.docx",
            "title": "ข้อกำหนดการใช้งาน (Terms of Service)"
        },
        {
            "md_file": "RANRHANนโยบายความเป็นส่วนตัวDRAFT.md",
            "out_docx": "RANRHAN_นโยบายความเป็นส่วนตัว_Privacy_Policy.docx",
            "title": "นโยบายความเป็นส่วนตัว (Privacy Policy)"
        },
        {
            "md_file": "RANRHANแนวทางแก้ไขแอปเพื่อรองรับกฎหมายDRAFT.md",
            "out_docx": "RANRHAN_แนวทางแก้ไขแอปเพื่อรองรับกฎหมาย_Roadmap.docx",
            "title": "แนวทางแก้ไขแอปเพื่อรองรับข้อกำหนดกฎหมาย"
        },
        {
            "md_file": "RANRHAN_สรุปรายการที่ต้องดำเนินการด้านกฎหมาย_Checklist.md",
            "out_docx": "RANRHAN_สรุปรายการที่ต้องดำเนินการด้านกฎหมาย_Checklist.docx",
            "title": "รายการตรวจสอบสิ่งที่ต้องดำเนินการด้านกฎหมาย"
        }
    ]
    
    print("=== Starting Conversion of Legal Documents to Word (.docx) ===")
    
    # 1. Convert each file individually
    for item in files_to_convert:
        md_path = os.path.join(base_dir, item["md_file"])
        docx_path = os.path.join(base_dir, item["out_docx"])
        
        if not os.path.exists(md_path):
            print(f"Warning: File not found: {md_path}")
            continue
            
        with open(md_path, 'r', encoding='utf-8') as f:
            content = f.read()
            
        doc = setup_document()
        parse_markdown_to_docx(content, doc)
        doc.save(docx_path)
        print(f"[OK] Generated: {item['out_docx']}")
        
    # 2. Create Master Combined Document
    combined_docx_path = os.path.join(base_dir, "RANRHAN_ชุดเอกสารกฎหมายและนโยบายความเป็นส่วนตัว_ฉบับสมบูรณ์.docx")
    master_doc = setup_document()
    
    # Add Cover Title for Combined Document
    cover_p = master_doc.add_paragraph()
    cover_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_p.paragraph_format.space_before = Pt(36)
    cover_p.paragraph_format.space_after = Pt(8)
    r_title = cover_p.add_run("ชุดเอกสารกฎหมายและนโยบายความเป็นส่วนตัว\nระบบ RAN-R-HAN")
    set_run_font(r_title, size_pt=24, bold=True, color=COLOR_TITLE)
    
    sub_p = master_doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(4)
    sub_p.paragraph_format.space_after = Pt(24)
    r_sub = sub_p.add_run("รวบรวม: ข้อกำหนดการใช้งาน (ToS) • นโยบายความเป็นส่วนตัว (PDPA) • แผนพัฒนาทางเทคนิค\nสถานะ: ฉบับร่างสำหรับทนายความและนักบัญชีตรวจสอบ (DRAFT)")
    set_run_font(r_sub, size_pt=14, italic=True, color=COLOR_MUTED)
    
    # Divider
    div_p = master_doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(20)
    p_bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1E3A8A"/></w:pBdr>')
    div_p._p.get_or_add_pPr().append(p_bdr)
    
    for idx, item in enumerate(files_to_convert):
        md_path = os.path.join(base_dir, item["md_file"])
        if os.path.exists(md_path):
            if idx > 0:
                master_doc.add_page_break()
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
            parse_markdown_to_docx(content, master_doc, is_first_doc=False)
            
    master_doc.save(combined_docx_path)
    print(f"[OK] Generated Master Document: RANRHAN_ชุดเอกสารกฎหมายและนโยบายความเป็นส่วนตัว_ฉบับสมบูรณ์.docx")
    print("=== Conversion Completed Successfully! ===")

if __name__ == "__main__":
    main()
